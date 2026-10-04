from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

SOURCE = Path(r"C:\Users\mesfi\Downloads\Telegram Desktop\Document1.docx")
OUTPUT = Path(r"C:\Users\mesfi\Downloads\Telegram Desktop\Document1_updated.docx")


def shade(cell, fill="D9E2F3"):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def table(document, headers, rows):
    t = document.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(t.rows[0].cells, headers):
        cell.text = text
        shade(cell)
        cell.paragraphs[0].runs[0].bold = True
        cell.paragraphs[0].runs[0].font.name = "Cambria"
        cell.paragraphs[0].runs[0].font.size = Pt(10)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row in rows:
        cells = t.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text
            cell.paragraphs[0].runs[0].font.name = "Cambria"
            cell.paragraphs[0].runs[0].font.size = Pt(10)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    document.add_paragraph()


def heading(document, text, level=1, page_break=False):
    p = document.add_paragraph(style=f"Heading {level}")
    if page_break:
        p.paragraph_format.page_break_before = True
    p.add_run(text)
    return p


def body(document, text, bullet=False):
    p = document.add_paragraph(style="Normal")
    if bullet:
        p.paragraph_format.left_indent = Inches(0.35)
        p.paragraph_format.first_line_indent = Inches(-0.2)
        p.add_run("- ")
    p.add_run(text)
    return p


def configure(document):
    n = document.styles["Normal"]
    n.font.name = "Cambria"
    n._element.rPr.rFonts.set(qn("w:eastAsia"), "Cambria")
    n.font.size = Pt(12)
    n.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    n.paragraph_format.line_spacing = 1.5
    n.paragraph_format.space_after = Pt(6)
    n.paragraph_format.first_line_indent = Inches(0.25)
    for level in range(1, 4):
        s = document.styles[f"Heading {level}"]
        s.font.name = "Cambria"
        s._element.rPr.rFonts.set(qn("w:eastAsia"), "Cambria")
        s.font.bold = True
        s.font.size = Pt(16 if level == 1 else 14 if level == 2 else 12)
        s.paragraph_format.space_before = Pt(12)
        s.paragraph_format.space_after = Pt(6)
        s.paragraph_format.keep_with_next = True


def page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def toc_field(paragraph):
    paragraph.text = ""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "Right-click and update field to generate the table of contents."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(text)
    run._r.append(end)


def insert_toc(document):
    anchor = next((p for p in document.paragraphs if p.text.strip() == "1. Introduction"), None)
    if anchor is None:
        return
    title = document.add_paragraph(style="Heading 1")
    title.paragraph_format.page_break_before = True
    title.add_run("Table of Contents")
    field = document.add_paragraph(style="Normal")
    toc_field(field)
    anchor._p.addprevious(title._p)
    anchor._p.addprevious(field._p)


def normalize(document):
    replacements = {
        "Our group worked in the [Data collection from all internet layer from surface ,deep and dark sites] Section of the organization. The exact name should be replaced with the department stated in our internship placement letter.": "Our group worked in the data-collection and analysis work area concerned with surface, deep, and authorized dark-web research. The formal department name should be completed from the internship placement letter before submission.",
        "during the period of [08/10/2018] to [02/13/2018]": "during the internship period recorded in the placement documentation",
        "Internee: XXXXXXXXXXXXXXX": "Internee: [Student name(s) and identification number(s) to be completed]",
        "(ID)": "(Student ID: [to be completed])",
        "Mentor: XXXXXXXXXXXX": "Mentor: [University mentor name to be completed]",
        "Company Supervisor: XXXXXXXXX": "Company Supervisor: Mr Tibebe",
    }
    for p in document.paragraphs:
        if p.text in replacements:
            p.text = replacements[p.text]
        p.text = p.text.replace("[08/10/2018] to [02/13/2018]", "the internship period recorded in the placement documentation")
        value = p.text.strip()
        if value.lower() in {"table of content", "table of contents"}:
            p.style = "Heading 1"
            p.paragraph_format.page_break_before = True
            toc_field(p)
            continue
        if value in {"Declaration", "Acknowledgements", "Executive Summary", "Internship Experience"}:
            p.style = "Heading 1"
            p.paragraph_format.page_break_before = True
        elif value == "Introduction":
            p.style = "Heading 2"
        elif value.startswith("Part "):
            p.style = "Heading 1"
        elif value and value[0].isdigit() and "." in value.split(" ", 1)[0]:
            p.style = "Heading 1" if value.split(" ", 1)[0].count(".") == 1 else "Heading 2"
        elif value in {"Debre Berhan University", "College of Computing", "Internship Report"}:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.bold = True
                run.font.name = "Cambria"
                run.font.size = Pt(16 if value == "Internship Report" else 14)


def add_project_work(document):
    heading(document, "Part Four: Project Work", 1, True)
    body(document, "This part presents DukaScraper as an implemented software system. The description is derived from the repository structure, application routes, worker modules, frontend pages, configuration files, tests, and operational documentation available with the project. It distinguishes implemented behavior from future recommendations and does not claim measurements or integrations that are not represented in the codebase.")
    heading(document, "4.1 Summary of the Project", 2)
    body(document, "DukaScraper is a distributed crawling and analytics pipeline for Amharic and English web content. Its current implementation is centered on Kafka-driven worker execution, recursive crawling, Amharic-aware parsing, and hosted Groq-based intelligence classification. The system is presented through a web interface and a FastAPI backend. A user can authenticate, submit a crawl job, follow the job state, inspect collected articles, review alerts and analytics, search stored information, and export results through the available application surfaces.")
    body(document, "The project is not a single crawler process. It is a coordinated set of services with separate responsibilities. Discovery and crawling are handled by workers specialized for the surface web, deeper JavaScript-enabled pages, and authorized dark-web connectivity. Parser and intelligence workers transform collected material into structured results. Persistence and analytics use different storage technologies for different access patterns. The UI communicates with API routes and can receive real-time updates through the project WebSocket support.")
    body(document, "The architecture exposes the complete software life cycle: requirements analysis, interface design, asynchronous processing, data modeling, error handling, security controls, testing, deployment, and monitoring. It also demonstrates the practical limitations of automated collection. A protected page is not treated as an invitation to bypass controls; instead, the system must record the result and communicate that authorized access is required.")

    heading(document, "4.2 Problem Statement and Justification", 2)
    body(document, "Organizations that need to monitor online information face a volume and speed problem. Websites change continuously, relevant material is distributed across different domains, and important evidence may appear in either English or Amharic. Manual collection is slow, difficult to repeat, and difficult to audit. A monitoring team also needs more than downloaded pages: it needs normalized text, extracted entities, duplicate control, severity classification, searchable storage, alerts, and a way to understand activity over time.")
    body(document, "The existing problem is an end-to-end information-processing problem. A solution must accept a collection request, schedule work, crawl within configured limits, parse content, detect language and quality, identify useful entities, classify intelligence, store the result, and make the result available to the requesting user. Each stage can fail independently because of network conditions, malformed pages, queue delays, or external access controls. The system must make these states visible instead of presenting incomplete data as if it were complete.")
    body(document, "DukaScraper is justified because its repository implements the foundation for this workflow. The API exposes job, article, storage, alert, analytics, monitoring, authentication, credential, and export routes. The workers separate collection from parsing and intelligence processing. The frontend contains dedicated pages for jobs, new crawls, job details, articles, search, storage, monitoring, analytics, and alerts. This division reduces the coupling between a user request and the long-running processing needed to fulfill it.")

    heading(document, "4.3 General and Specific Objectives", 2)
    body(document, "The general objective is to implement a distributed web-crawling and threat-analytics system that can collect authorized online content, process English and Amharic material, preserve structured evidence, and present useful results through a secure web interface.")
    for item in [
        "Provide authenticated users with a controlled way to submit and monitor crawl jobs.",
        "Separate collection, discovery, parsing, intelligence, export, and health responsibilities into maintainable worker and service components.",
        "Support surface-web crawling and specialized processing paths for deeper JavaScript content and authorized dark-web connectivity.",
        "Normalize and analyze English and Amharic web text using the language modules present in the repository.",
        "Extract entities and classify content using hosted LLM intelligence when available and an enhanced fallback path when it is not.",
        "Store structured records, raw or processed objects, operational data, and analytical information using appropriate persistence services.",
        "Provide alerts, analytics, search, article inspection, exports, and monitoring views through the implemented API and React frontend.",
        "Protect user-owned data through authentication, authorization, rate limiting, scoped access, and user-specific filtering.",
        "Make operational failures diagnosable through logs, metrics, health checks, retries, and explicit status reporting.",
    ]:
        body(document, item, True)

    heading(document, "4.4 Implemented System Scope", 2)
    table(document, ["Subsystem", "Verified implementation surface", "Purpose"], [
        ("API", "app/api/routes", "Authentication, jobs, articles, alerts, analytics, storage, credentials, monitoring, and exports."),
        ("Pipeline", "app/pipeline", "Schemas, producers, consumers, topics, and asynchronous processing."),
        ("Surface worker", "workers/surface-worker", "Collection of ordinary surface-web content through the configured pipeline."),
        ("Deep worker", "workers/deep-worker", "Processing pages requiring deeper or browser-capable collection behavior."),
        ("Dark worker", "workers/dark-worker", "Separate path for configured, authorized dark-web connectivity."),
        ("Parser and LLM workers", "workers/parser-worker and workers/llm-worker", "Content parsing, entity extraction, and intelligence classification."),
        ("Frontend", "ui/src/pages and ui/src/components", "React and TypeScript experience for jobs, results, alerts, analytics, search, storage, and monitoring."),
        ("Operations", "Docker Compose, monitoring, health, and metrics modules", "Repeatable local deployment, readiness, metrics, and diagnostics."),
    ])

    heading(document, "4.5 System Architecture", 2)
    body(document, "The system follows a service-oriented and event-driven arrangement. The user interface sends a request to the FastAPI backend. The backend validates the authenticated user and request, creates or updates job state, and publishes work through the pipeline. Kafka provides the event boundary between request handling and worker execution. Workers consume the relevant topic, perform their stage, and publish the next event or a failure state. A slow crawl therefore does not need to block the HTTP request until every page has been processed.")
    body(document, "The architecture creates explicit ownership boundaries. The API owns request validation and user-facing access control. Producers and consumers own message movement. Crawling workers own network collection. The parser owns extraction of useful document content. The intelligence worker owns entity and category analysis. Storage and export components own persistence and delivery. The frontend owns presentation and interaction, while monitoring surfaces expose service condition and operational measurements.")
    body(document, "The logical sequence is: User interface -> FastAPI API -> Kafka job topic -> discovery or crawling worker -> parser worker -> intelligence worker -> PostgreSQL, ClickHouse, Redis, or MinIO as appropriate -> alerts, analytics, article views, search, and exports.")
    table(document, ["Stage", "Input", "Processing responsibility", "Output"], [
        ("Request", "Authenticated crawl configuration", "Validate scope, ownership, and request fields.", "Job record and queued work."),
        ("Discovery", "Seed URL and crawl constraints", "Find links and schedule reachable pages.", "Discovered targets and crawl events."),
        ("Collection", "A permitted target", "Fetch content using the configured worker path.", "Raw response or explicit failure state."),
        ("Parsing", "Collected response", "Extract readable content, metadata, and article-level structure.", "Normalized document record."),
        ("Language and quality", "Normalized text", "Detect language and assess usable content quality.", "Language, quality, and processing metadata."),
        ("Intelligence", "Processed text", "Extract entities and categories through LLM or fallback analysis.", "Structured intelligence result."),
        ("Persistence", "Structured result", "Store records, objects, analytical facts, and cache state.", "Queryable user-owned result."),
        ("Presentation", "Stored result and events", "Display status, articles, alerts, analytics, and exports.", "User-visible monitoring and evidence."),
    ])

    heading(document, "4.6 Data Flow and Event Processing", 2)
    body(document, "An event-driven pipeline is useful because collection time is unpredictable. A small page may complete quickly, while a site with many links, JavaScript requirements, or a slow response can take much longer. The event boundary allows workers to process independent items, supports retries at appropriate stages, and makes the current state observable. The repository contains producer, consumer, topic, and schema modules that provide the code-level foundation for this behavior.")
    body(document, "A crawl job should be understood as a stateful unit rather than a single network call. Its record can move through creation, queueing, running, completion, partial completion, or failure states. Individual articles may have their own processing status. This separation matters because a crawl can produce useful results even when one page fails. The frontend job and job-detail pages exist to expose that distinction to users.")
    body(document, "The API includes alert routes, the frontend includes an alert page and alert bell component, and the project contains real-time support through a WebSocket hook and websocket package. These surfaces allow a user to notice important results without repeatedly reloading every page. The report does not claim that a particular deployment has delivered a specific number of alerts; it records the implemented feature path.")

    heading(document, "4.7 Crawling and Collection Workers", 2)
    body(document, "The surface worker is intended for ordinary web collection. The deep worker handles pages that need more involved processing, such as JavaScript-rendered content or deeper traversal behavior represented by the project configuration. The dark worker is isolated because dark-web connectivity has different operational, legal, and security requirements. Separate workers reduce the risk that special behavior required by one network path will be silently applied to all requests.")
    body(document, "The worker separation improves deployment and diagnosis. A team can inspect the health, logs, configuration, and resource use of a worker independently. The repository includes shared worker utilities, health support, metrics support, Dockerfiles, startup scripts, and readiness-related files. These are implementation mechanisms, not merely architectural aspirations.")
    body(document, "External access restrictions are a central limitation. Web owners may use robots rules, CAPTCHAs, rate limits, authentication, browser verification, or other controls. The system should not be described as defeating these controls. Its responsible behavior is to stop or classify the condition, record a useful diagnostic state, and require authorization where access is restricted.")

    heading(document, "4.8 Parsing, Language Processing, and Quality", 2)
    body(document, "The repository contains language packages for cleaning, deduplication, language detection, normalization, quality assessment, and tokenization. These modules reflect the fact that collection alone is not analysis. Web pages contain navigation text, repeated headers, advertisements, malformed markup, duplicate URLs, and mixed-language material. Cleaning and normalization make later classification more consistent, while quality checks reduce the risk of treating an empty or mostly boilerplate page as meaningful evidence.")
    body(document, "Amharic support is a defining project concern. A useful Amharic-aware pipeline must preserve the Ethiopic script, avoid applying English-only assumptions, and retain enough normalized text for search and intelligence extraction. English support remains important because many sites and technical indicators use English terms, domain names, email addresses, and security vocabulary. The two language contexts can appear in the same monitoring task, so detection and normalization are processing metadata rather than a hard assumption about the whole system.")
    body(document, "Deduplication is similarly important. The same story can appear at multiple URLs, through tracking parameters, or in slightly different copies. Duplicate control reduces storage waste, repeated alerts, and misleading analytics. The project includes deduplication and content-fingerprint services, as well as Redis-related support for temporary state and duplicate control.")

    heading(document, "4.9 Entity Extraction and Intelligence Classification", 2)
    body(document, "The intelligence stage turns processed text into structured information. The LLM worker includes hosted intelligence processing and an enhanced fallback analysis path. The fallback path matters because external model calls can fail through missing credentials, timeouts, response parsing errors, or temporary service availability. A resilient system should make the degraded path visible in logs and preserve the result shape expected by downstream consumers.")
    body(document, "The implemented improvement work documented for the repository includes three attempts with exponential backoff, smart truncation that preserves the head, sampled body, and tail of long content, and expanded fallback patterns for email addresses, IP addresses, domains, proper nouns, and organization names. The fallback result is deduplicated case-insensitively, filtered for short noise, and capped. The intelligence prompt asks for all named entities and identifies expected entity types. These are concrete processing choices and do not claim that every model response is perfect.")
    body(document, "Classification is useful for prioritization rather than as a replacement for human judgment. Category and severity results help an analyst decide what to review first, but a classification may be affected by language, missing context, truncation, or an unavailable model. Logging of retries, parsing failures, fallback use, and entity counts is therefore part of the feature.")

    heading(document, "4.10 Storage, Search, and Export", 2)
    body(document, "The project uses multiple storage technologies because the data has different shapes and access patterns. PostgreSQL is appropriate for relational records such as users, jobs, permissions, alerts, and structured processing state. ClickHouse is suited to analytical queries over larger event or result sets. Redis supports temporary state, cache behavior, rate control, and duplicate detection. MinIO provides object storage for raw crawl content or processed files. The report identifies these roles without assuming that every record is written to every store.")
    body(document, "The API route set includes articles, storage, analytics, search-related frontend behavior, and exports. This indicates that the project is designed for both operational monitoring and later review. A result is more valuable when a user can inspect its source context, search related content, observe trends, and export selected information. Data ownership remains a security requirement: normal users must receive only records permitted by their authenticated identity and scope.")
    table(document, ["Need", "Relevant implementation area", "Expected user value"], [
        ("Operational records", "PostgreSQL-backed service and API routes", "Consistent jobs, users, alerts, and ownership records."),
        ("Large-scale analysis", "ClickHouse configuration and analytics route", "Aggregations and trend-oriented views."),
        ("Temporary state", "Redis configuration and services", "Fast coordination, caching, throttling, and duplicate checks."),
        ("Raw or file content", "MinIO and storage route", "Access to stored objects without overloading relational rows."),
        ("User retrieval", "Articles, search, storage, and export routes", "Review and reuse of collected evidence."),
    ])

    heading(document, "4.11 Security and Privacy Controls", 2)
    body(document, "The application includes authentication routes and security modules for authentication, rate limiting, and scope. The frontend contains login, email verification, one-time-password verification, forgot-password, and change-password pages. These surfaces represent a complete account workflow rather than a dashboard that assumes every visitor is trusted. Credentials and authentication state must be handled through the configured environment and should not be placed in source code or report screenshots.")
    body(document, "Authorization is equally important. A user who submits a crawl should not automatically be able to see another user's jobs, articles, alerts, stored content, searches, or analytics. The internship work identified user-data isolation as a problem and updated queries and access paths to filter by the authenticated user identifier. This is a behavioral requirement that should be tested at the API boundary, not only assumed from a frontend filter.")
    body(document, "Rate limiting protects both the service and external targets from uncontrolled request volume. Scope checks reduce the chance that a route is called outside the permissions granted to the current identity. Operational logs should avoid exposing passwords, tokens, or sensitive collected content unnecessarily. Because DukaScraper concerns online information and potentially sensitive indicators, authorization, data minimization, retention, and responsible disclosure should remain part of system maintenance.")

    heading(document, "4.12 Frontend and User Interaction", 2)
    body(document, "The frontend is implemented with React, TypeScript, and Vite. Its pages cover the principal user workflow: login and verification, dashboard overview, creation of a new crawl, job list, job detail, articles, search, storage, alerts, analytics, monitoring, and user administration. Reusable components include layout, logo, alert bell, health pill, performance panel, selectors, OTP input, and shared UI elements. This organization supports consistent navigation and makes page-level behavior easier to test and maintain.")
    body(document, "The dashboard is not only a decorative view. It gives the user a starting point for current job state, operational health, and recent results. Job pages provide access to work history. Article and storage pages support evidence review. Analytics summarizes patterns, while alerts direct attention to higher-priority events. Monitoring helps distinguish a data problem from a service problem. Together these views form a practical analyst workflow from request to review.")
    body(document, "The project includes automatic refresh and real-time support hooks. Refresh behavior must be bounded and should respect page visibility, server load, and authentication state. WebSocket updates must handle disconnects without freezing the rest of the interface. These operational details affect user trust: a dashboard should communicate whether information is current, delayed, incomplete, or unavailable.")

    heading(document, "4.13 Development Methodology", 2)
    body(document, "The project was developed iteratively. The group first identified a concrete problem, inspected the existing code, reproduced or localized the behavior, selected the smallest responsible component, implemented a focused change, and tested the result. Logs and health checks were used together with unit tests and API checks. This method was especially useful in a distributed system where an apparent frontend issue can be caused by a worker, queue, database query, or environment variable.")
    body(document, "The repository uses Python application code, worker modules, TypeScript frontend code, Docker configuration, test suites, and operational scripts. Each tool has a different role. Git records changes, pytest exercises behavior, Docker Compose supplies local service topology, and monitoring tools expose runtime condition. The development method is both software engineering and systems engineering: a change is complete only when the component works and its integration behavior remains understandable.")
    body(document, "A practical review loop is: define expected behavior; identify the owning route, worker, or service; reproduce the issue with the narrowest test; change the controlling logic; run the focused test; inspect logs and related integration behavior; and document the result. This reduces speculative edits and makes it easier for another developer to continue the work.")

    heading(document, "4.14 Results and Discussion", 2)
    body(document, "The implemented result is a project structure for distributed collection and analysis rather than a claim that every possible website can be crawled. The repository contains services, workers, routes, frontend pages, language modules, storage configuration, and test organization required for the described workflow. The result is strongest as an extensible monitoring foundation: it separates concerns, provides user-facing review paths, and leaves clear places for future hardening.")
    body(document, "The entity-extraction improvements demonstrate the value of measuring failure modes. Before the improvement, a missing key, truncation, parse error, or timeout could result in an empty extraction or weak regex output. The updated logic retries transient failures, distributes long-content sampling, improves fallback patterns, and logs the reason for degradation. The documented test set covers proper nouns and organizations, technical entities, case-insensitive deduplication, category detection, truncation behavior, missing-key logging, and comparison of fallback quality.")
    body(document, "The result should still be evaluated with real deployment data. Useful measurements include successful versus failed jobs, median processing time, queue delay, parser quality, duplicate rate, model fallback frequency, alert precision, storage growth, and per-worker error rates. These are recommendations for operational evaluation, not completed benchmarks, because the repository context does not supply a verified production dataset or performance report.")

    heading(document, "4.15 Project Screenshots and Evidence", 2)
    body(document, "The final submission should include screenshots captured from the running project rather than generic illustrations. The evidence list is tied to implemented frontend pages and should be populated from the configured local deployment: login and verification, dashboard, new crawl form, jobs list, job detail, article review, search, storage, alerts, analytics, monitoring, and user management. Each screenshot should show the date, environment, and a short caption. Credentials, tokens, private URLs, and sensitive collected content must be removed before insertion.")
    table(document, ["Figure", "Evidence to capture", "Purpose"], [
        ("Figure 1", "Authenticated dashboard", "Show the main operational entry point."),
        ("Figure 2", "New crawl form", "Show how an authorized user starts collection."),
        ("Figure 3", "Jobs and job detail", "Show state tracking and processing history."),
        ("Figure 4", "Article or search view", "Show normalized content review."),
        ("Figure 5", "Alerts and analytics", "Show prioritization and aggregate results."),
        ("Figure 6", "Monitoring page", "Show service-health visibility."),
    ])

    heading(document, "4.16 Limitations and Responsible Use", 2)
    body(document, "The system depends on external websites, network conditions, configured services, and optional hosted intelligence. It cannot guarantee availability or completeness of external content. A page may change after collection, refuse automated access, require a login, or return content that is not suitable for parsing. The user interface and reports should preserve these distinctions so that analysts do not confuse absence of evidence with evidence of absence.")
    body(document, "The system must be used only for authorized collection. Dark-web connectivity, browser automation, credential handling, and threat-related data require stronger controls than an ordinary public-site reader. Operators should follow organizational policy, applicable Ethiopian law, site terms, privacy requirements, and explicit authorization. The design choice to report access restrictions rather than bypass them is a project strength and should remain a maintenance rule.")

    heading(document, "4.17 Conclusion and Recommendation", 2)
    body(document, "DukaScraper demonstrates how a distributed application can connect crawling, language processing, intelligence extraction, storage, alerts, analytics, and a web interface. Its strongest engineering qualities are separation of responsibilities, explicit worker paths, multiple storage choices, user-facing review surfaces, and attention to fallback behavior. Its next improvements should focus on measurable operational reliability, security review, data-retention policy, controlled authorized integrations, and additional tests around cross-user isolation and partial job failure.")
    body(document, "The project should be maintained through small, testable changes. New features should identify their owning route or worker, define the expected data contract, add a focused test, update monitoring or logs, and then update the documentation. This practice will help the system grow without losing the clarity that makes a distributed architecture manageable.")


def appendices(document):
    heading(document, "Part Five: General Conclusion and Recommendation", 1, True)
    body(document, "The internship connected university knowledge with the practical construction of a distributed software system. The work required more than writing code: it required reading an existing repository, tracing data across services, protecting user records, respecting external access controls, testing failure paths, and communicating technical results. DukaScraper provided a realistic setting in which backend, frontend, database, messaging, security, and operations concerns had to be considered together.")
    body(document, "The recommended direction is controlled improvement. The organization should continue giving interns real systems with clear supervision and safe test environments. The project team should maintain a verified feature inventory, add regression tests for every security boundary, measure worker reliability, document deployment prerequisites, and capture screenshots only from an authorized running environment. The system should not expand collection capability by weakening authorization or attempting to defeat anti-bot protection.")

    heading(document, "References", 1, True)
    for item in [
        "DukaScraper repository source code, configuration, tests, and documentation, accessed during report preparation.",
        "DukaScraper README.md: project overview and supported processing scope.",
        "DukaScraper application routes under app/api/routes: implemented API surface.",
        "DukaScraper pipeline modules under app/pipeline: schemas, producers, consumers, and topics.",
        "DukaScraper language modules under app/language: cleaning, detection, normalization, quality, tokenization, and deduplication.",
        "DukaScraper worker modules under workers: surface, deep, dark, discovery, parser, LLM, exporter, and transcription workers.",
        "DukaScraper frontend source under ui/src: React pages, components, API client, real-time hooks, and types.",
        "DukaScraper test suites under tests: unit, integration, system, and performance test organization.",
        "Internship Report format-2017.docx, Debre Berhan University, supplied university format document.",
        "Information Network Security Administration internship materials supplied by the internship host or university, where applicable.",
    ]:
        body(document, item, True)

    heading(document, "Appendix A: Implemented Feature Verification Matrix", 1, True)
    table(document, ["Feature", "Repository evidence", "Report treatment"], [
        ("Authentication and recovery", "auth.py; login, verification, OTP, forgot-password, and change-password pages", "Described as implemented account workflow."),
        ("Crawl management", "jobs.py; NewCrawl, Jobs, and JobDetail pages", "Described as job submission and tracking."),
        ("Articles and storage", "articles.py and storage.py; Article, Storage, and JobArticles pages", "Described as result inspection and stored-content access."),
        ("Alerts and analytics", "alerts.py and analytics.py; Alerts and Analytics pages", "Described as implemented review surfaces."),
        ("Monitoring", "monitoring routes, proxy, health, and metrics modules", "Described as operational visibility."),
        ("Kafka pipeline", "pipeline producers, consumers, topics, and schemas", "Described as event-driven processing."),
        ("Specialized workers", "surface, deep, dark, discovery, parser, LLM, exporter, and transcribe workers", "Described with authorization limitations."),
        ("Language processing", "app/language package and parser worker", "Described as English and Amharic-aware processing."),
        ("Intelligence fallback", "LLM worker and documented improvement tests", "Described with retries, truncation, fallback, and logging."),
        ("User isolation", "security scope/auth modules and route filtering work", "Described as a privacy behavior."),
    ])

    heading(document, "Appendix B: Suggested Test and Acceptance Checklist", 1, True)
    for item in [
        "Create a test account and verify the configured authentication, email, and OTP workflow.",
        "Submit a crawl job and confirm that it appears only in the owning account's job list.",
        "Confirm that a second normal user cannot read the first user's articles, alerts, storage, searches, or analytics.",
        "Process a short English document and verify parsing, language metadata, entity extraction, and persistence.",
        "Process a short Amharic document and verify that Ethiopic text is retained through normalization and display.",
        "Process a long document and verify that smart truncation preserves head, sampled body, tail, and markers.",
        "Disable the hosted model credential in a test environment and verify fallback analysis and diagnostics.",
        "Simulate a temporary model failure and verify retry behavior and final fallback or failure status.",
        "Submit duplicate content and verify content-fingerprint or duplicate-control behavior.",
        "Observe a worker health endpoint and confirm that monitoring reflects unavailable or recovering services.",
        "Export a permitted result and confirm that the export contains only authorized data.",
        "Record version, configuration, and date for each screenshot inserted into the final report.",
    ]:
        body(document, item, True)

    heading(document, "Appendix C: Operational Configuration Notes", 1, True)
    body(document, "The repository includes Docker and Docker Compose configuration, monitoring configuration, database directories, worker startup scripts, and environment-dependent service settings. Before a demonstration, operators should verify that required services are healthy, Kafka topics are available, PostgreSQL and analytical storage are reachable, Redis and object storage are configured, and the frontend points to the intended API. Secrets should be supplied through the deployment environment and excluded from screenshots and report text.")
    body(document, "A useful troubleshooting order is to begin at the user-visible symptom, identify the API request, inspect the job or record, check the relevant Kafka event and worker log, confirm downstream storage, and then return to the frontend. This prevents a presentation problem from being confused with a processing problem. For a missing alert, inspect alert creation, severity classification, ownership filtering, persistence, and frontend retrieval separately.")
    body(document, "Operational records should be retained only as long as the organization requires. Raw content, extracted entities, credentials, logs, and exports may have different sensitivity and retention needs. A future production policy should specify who may access each class of data, how long it is stored, how it is deleted, and how an incident is reported.")

    heading(document, "Appendix D: Glossary", 1, True)
    table(document, ["Term", "Meaning in this project"], [
        ("API", "Application Programming Interface used by the frontend and other clients."),
        ("Amharic-aware processing", "Text handling that preserves and analyzes content written in Ethiopic script."),
        ("Consumer", "A pipeline component that receives an event from a Kafka topic."),
        ("Crawler", "A component that retrieves permitted web resources according to configured rules."),
        ("Entity", "A named or technical item such as a person, organization, location, domain, IP, or email."),
        ("Fallback", "A controlled alternative used when the preferred processing path is unavailable or fails."),
        ("Kafka", "Event-streaming and message-queue technology used to connect pipeline stages."),
        ("MinIO", "Object storage used for files or raw and processed content."),
        ("Parser", "A component that converts collected responses into usable structured content."),
        ("Producer", "A pipeline component that publishes an event to a Kafka topic."),
        ("Redis", "Fast temporary storage used for cache, coordination, rate control, or duplicate checks."),
        ("Scope", "The permission boundary determining which resources an identity may use."),
        ("Severity", "A prioritization value assigned to an analyzed result or alert."),
        ("Worker", "An independently deployable process responsible for one processing stage."),
    ])


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    document = Document(str(SOURCE))
    configure(document)
    normalize(document)
    insert_toc(document)
    add_project_work(document)
    appendices(document)
    section = document.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.25)
    section.right_margin = Inches(1.0)
    page_number(section.footer.paragraphs[0])
    document.core_properties.title = "DukaScraper Internship Report"
    document.core_properties.subject = "Verified implementation and internship experience report"
    document.save(str(OUTPUT))
    print(f"WROTE {OUTPUT}")
    print(f"PARAGRAPHS {len(document.paragraphs)} TABLES {len(document.tables)}")
    print(f"WORDS {sum(len(p.text.split()) for p in document.paragraphs)}")


if __name__ == "__main__":
    main()
